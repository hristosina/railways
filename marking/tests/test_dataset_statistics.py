import tempfile
import unittest
from pathlib import Path

import yaml
from openpyxl import load_workbook

from dataset_statistics import (
    collect_dataset_statistics,
    discover_dataset_scenarios,
    generate_dataset_statistics_report,
)


class DatasetStatisticsTests(unittest.TestCase):
    def make_dataset(self, root):
        root = Path(root)
        for split in ("train", "valid", "test"):
            (root / split / "images").mkdir(parents=True)
            (root / split / "labels").mkdir(parents=True)
        (root / "data.yaml").write_text(
            yaml.safe_dump({
                "train": "train/images",
                "val": "valid/images",
                "test": "test/images",
                "names": ["vehicle", "rails"],
            }, allow_unicode=True, sort_keys=False),
            encoding="utf-8",
        )

        def image(split, name, labels=None):
            (root / split / "images" / name).write_bytes(b"image")
            if labels is not None:
                (root / split / "labels" / f"{Path(name).stem}.txt").write_text(
                    labels, encoding="utf-8"
                )

        image("train", "a.jpg", "0 0.5 0.5 0.2 0.2\n0 0.4 0.4 0.1 0.1\n1 0.5 0.5 0.3 0.3\n")
        image("train", "without_label.jpg")
        image("valid", "empty.jpg", "")
        image("test", "test_afternoon_001.jpg", "0 0.5 0.5 0.2 0.2\n")
        image("test", "test_night_001.jpg", "1 0.5 0.5 0.2 0.2\n")
        image("test", "unknown_001.jpg", "0 0.5 0.5 0.2 0.2\n")
        return root

    def test_collects_splits_classes_scenarios_and_problems(self):
        with tempfile.TemporaryDirectory() as directory:
            root = self.make_dataset(directory)
            stats = collect_dataset_statistics(root)

            self.assertEqual(stats["splits"]["train"]["image_count"], 2)
            self.assertEqual(stats["splits"]["train"]["instances"], [2, 1])
            self.assertEqual(stats["splits"]["train"]["images_with_class"], [1, 1])
            self.assertEqual(stats["splits"]["train"]["missing_labels"], 1)
            self.assertEqual(stats["splits"]["val"]["empty_annotations"], 1)
            self.assertEqual(stats["splits"]["test"]["image_count"], 3)
            self.assertEqual(
                [scenario["name"] for scenario in stats["scenarios"]],
                ["День", "Ночь", "unknown"],
            )
            self.assertEqual(len(stats["problems"]), 1)

    def test_excel_contains_all_required_sheets_and_native_charts(self):
        with tempfile.TemporaryDirectory() as directory:
            root = self.make_dataset(directory)
            output = root / "Статистика_датасета.xlsx"
            stats, path = generate_dataset_statistics_report(root, output)

            self.assertEqual(path, output.resolve())
            workbook = load_workbook(path, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                ["Сводка", "Train", "Val", "Test", "Сценарии test", "Проблемы"],
            )
            self.assertEqual(workbook["Train"]["B4"].value, 2)
            self.assertEqual(workbook["Train"]["D12"].value, 2)
            self.assertEqual(len(workbook["Train"]._charts), 1)
            self.assertEqual(len(workbook["Val"]._charts), 1)
            self.assertEqual(len(workbook["Test"]._charts), 1)
            self.assertEqual(len(workbook["Сценарии test"]._charts), 1)
            train_chart = workbook["Train"]._charts[0]
            scenario_chart = workbook["Сценарии test"]._charts[0]
            self.assertIsNone(train_chart.dLbls)
            self.assertIsNone(scenario_chart.dLbls)
            self.assertEqual(len(train_chart.series[0].dPt), 2)
            self.assertTrue(train_chart.varyColors)
            self.assertFalse(train_chart.y_axis.delete)
            self.assertEqual(train_chart.y_axis.tickLblPos, "nextTo")
            self.assertIn("Количество экземпляров", str(train_chart.y_axis.title))
            self.assertEqual(scenario_chart.x_axis.tickLblSkip, 1)
            self.assertIsNone(scenario_chart.x_axis.title)
            self.assertIsNone(train_chart.legend)
            self.assertFalse(train_chart.y_axis.title.tx.rich.p[0].r[0].rPr.b)
            self.assertAlmostEqual(train_chart.layout.manualLayout.x, 0.07)
            self.assertAlmostEqual(train_chart.layout.manualLayout.w, 0.82)
            self.assertFalse(scenario_chart.legend.overlay)
            self.assertEqual(workbook["Сценарии test"]["C6"].value, "Транспорт")
            self.assertEqual(workbook["Сценарии test"]["D6"].value, "Рельсы")
            self.assertIn(stats["dataset_root"], workbook["Сводка"]["B3"].value)

    def test_discovers_scenarios_from_actual_filename_prefixes(self):
        with tempfile.TemporaryDirectory() as directory:
            root = self.make_dataset(directory)
            definitions = discover_dataset_scenarios(root)

            self.assertEqual(
                [definition["name"] for definition in definitions],
                ["День", "Ночь", "unknown"],
            )
            self.assertEqual(
                [definition["prefix"] for definition in definitions],
                ["afternoon", "night", "unknown"],
            )

    def test_explicit_scenarios_can_be_renamed_reordered_and_filtered(self):
        with tempfile.TemporaryDirectory() as directory:
            root = self.make_dataset(directory)
            test_images = root / "test" / "images"
            definitions = [
                {"name": "Пользовательская ночь", "path": str(test_images), "prefix": "night"},
                {"name": "Пользовательский день", "path": str(test_images), "prefix": "afternoon"},
            ]

            stats = collect_dataset_statistics(root, scenarios=definitions)

            self.assertEqual(
                [scenario["name"] for scenario in stats["scenarios"]],
                ["Пользовательская ночь", "Пользовательский день"],
            )
            self.assertEqual(
                [scenario["image_count"] for scenario in stats["scenarios"]],
                [1, 1],
            )

    def test_empty_explicit_scenario_list_disables_scenario_breakdown(self):
        with tempfile.TemporaryDirectory() as directory:
            root = self.make_dataset(directory)
            stats = collect_dataset_statistics(root, scenarios=[])
            self.assertEqual(stats["scenarios"], [])


if __name__ == "__main__":
    unittest.main()

"""Сбор воспроизводимой статистики YOLO-датасета и экспорт в Excel."""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from dataset_utils import IMAGE_EXTENSIONS, find_dataset_yaml, normalize_dataset_path


SPLIT_LABELS = {"train": "Обучающая", "val": "Валидационная", "test": "Тестовая"}
SCENARIO_NAMES = {
    "fallout": "Осадки", "precipitation": "Осадки", "rain": "Осадки", "осадки": "Осадки",
    "fog": "Туман", "туман": "Туман",
    "afternoon": "День", "day": "День", "день": "День",
    "night": "Ночь", "ночь": "Ночь",
    "twilight": "Сумерки", "dusk": "Сумерки", "сумерки": "Сумерки",
}
SCENARIO_ORDER = {name: index for index, name in enumerate(("Осадки", "Туман", "День", "Сумерки", "Ночь"))}
CLASS_NAMES = {
    "vehicle": "Транспорт", "transport": "Транспорт", "транспорт": "Транспорт",
    "train": "Поезд", "поезд": "Поезд",
    "person": "Человек", "human": "Человек", "человек": "Человек",
    "rails": "Рельсы", "rail": "Рельсы", "рельсы": "Рельсы",
}
CLASS_ORDER = {name: index for index, name in enumerate(("Транспорт", "Поезд", "Человек", "Рельсы"))}
EXCEL_SERIES_COLORS = {
    "Транспорт": "4F81BD",
    "Поезд": "C0504D",
    "Человек": "9BBB59",
    "Рельсы": "8064A2",
}
SUMMARY_POINT_COLORS = ("4F81BD", "C0504D", "9BBB59")
SCENARIO_RE = re.compile(r"^(?:test_)?(?P<scenario>[^_]+)_", re.IGNORECASE)


def display_class(name):
    return CLASS_NAMES.get(str(name).strip().lower(), str(name))


def display_scenario(name):
    return SCENARIO_NAMES.get(str(name).strip().lower(), str(name))


def _class_names(data):
    names = data.get("names", [])
    if isinstance(names, dict):
        names = [names[key] for key in sorted(names, key=lambda value: int(value))]
    if not isinstance(names, list) or not names or not all(isinstance(name, str) for name in names):
        raise ValueError("В data.yaml не найден корректный список names.")
    return names


def _as_image_directory(path):
    path = Path(path)
    if path.name.lower() == "images" and path.is_dir():
        return path
    if (path / "images").is_dir():
        return path / "images"
    return path if path.is_dir() else None


def _resolve_split_directories(dataset_root, yaml_path, data, split):
    aliases = ("val", "valid") if split == "val" else (split,)
    values = []
    for key in aliases:
        configured = data.get(key)
        if configured:
            values.extend(configured if isinstance(configured, list) else [configured])

    yaml_base = yaml_path.parent
    configured_base = data.get("path")
    if configured_base:
        configured_base = Path(str(configured_base)).expanduser()
        if not configured_base.is_absolute():
            configured_base = yaml_base / configured_base
    bases = [configured_base, yaml_base, dataset_root]
    result = []
    for value in values:
        value_path = Path(str(value)).expanduser()
        candidates = [value_path] if value_path.is_absolute() else [base / value_path for base in bases if base]
        for candidate in candidates:
            images_dir = _as_image_directory(candidate.resolve())
            if images_dir and images_dir not in result:
                result.append(images_dir)
                break

    for alias in aliases:
        fallback = dataset_root / alias / "images"
        if fallback.is_dir() and fallback not in result:
            result.append(fallback)
    return result


def _image_label_pairs(images_dir):
    images_dir = Path(images_dir)
    labels_dir = images_dir.parent / "labels"
    pairs = []
    for image_path in sorted(images_dir.rglob("*"), key=lambda path: str(path).casefold()):
        if not image_path.is_file() or image_path.suffix.lower() not in IMAGE_EXTENSIONS:
            continue
        if ".marking_trash" in image_path.parts:
            continue
        relative = image_path.relative_to(images_dir).with_suffix(".txt")
        pairs.append((image_path, labels_dir / relative))
    return pairs


def _analyze_pairs(pairs, class_names, section_name):
    instances = [0] * len(class_names)
    images_with_class = [0] * len(class_names)
    annotated_images = 0
    empty_annotations = 0
    missing_labels = 0
    invalid_lines = 0
    problems = []

    for image_path, label_path in pairs:
        if not label_path.is_file():
            missing_labels += 1
            problems.append({
                "section": section_name,
                "image": str(image_path),
                "label": str(label_path),
                "problem": "Отсутствует файл разметки",
                "details": "",
            })
            continue

        seen_classes = set()
        valid_objects = 0
        try:
            lines = label_path.read_text(encoding="utf-8").splitlines()
        except (OSError, UnicodeError) as exc:
            invalid_lines += 1
            problems.append({
                "section": section_name,
                "image": str(image_path),
                "label": str(label_path),
                "problem": "Не удалось прочитать разметку",
                "details": str(exc),
            })
            continue

        for line_number, line in enumerate(lines, start=1):
            parts = line.split()
            try:
                class_id = int(parts[0])
                if len(parts) < 5 or not 0 <= class_id < len(class_names):
                    raise ValueError
                # Проверяем числовой формат координат, но не меняем исходную разметку.
                [float(value) for value in parts[1:]]
            except (IndexError, TypeError, ValueError):
                invalid_lines += 1
                problems.append({
                    "section": section_name,
                    "image": str(image_path),
                    "label": str(label_path),
                    "problem": "Некорректная строка разметки",
                    "details": f"Строка {line_number}: {line[:160]}",
                })
                continue
            instances[class_id] += 1
            seen_classes.add(class_id)
            valid_objects += 1

        if valid_objects:
            annotated_images += 1
            for class_id in seen_classes:
                images_with_class[class_id] += 1
        else:
            empty_annotations += 1

    return {
        "image_count": len(pairs),
        "annotated_images": annotated_images,
        "empty_annotations": empty_annotations,
        "missing_labels": missing_labels,
        "invalid_lines": invalid_lines,
        "instances": instances,
        "images_with_class": images_with_class,
        "problems": problems,
    }


def _merge_split_statistics(parts, paths, class_count):
    merged = {
        "paths": [str(path) for path in paths],
        "image_count": 0,
        "annotated_images": 0,
        "empty_annotations": 0,
        "missing_labels": 0,
        "invalid_lines": 0,
        "instances": [0] * class_count,
        "images_with_class": [0] * class_count,
        "problems": [],
    }
    for part in parts:
        for key in ("image_count", "annotated_images", "empty_annotations", "missing_labels", "invalid_lines"):
            merged[key] += part[key]
        merged["instances"] = [left + right for left, right in zip(merged["instances"], part["instances"])]
        merged["images_with_class"] = [
            left + right for left, right in zip(merged["images_with_class"], part["images_with_class"])
        ]
        merged["problems"].extend(part["problems"])
    return merged


def _scenario_key_from_file(filename):
    match = SCENARIO_RE.match(filename)
    return match.group("scenario").casefold() if match else None


def _auto_scenario_definitions(dataset_root, test_directories):
    """Находит сценарии по структуре папок или фактическим префиксам файлов."""
    physical = []
    search_roots = (dataset_root / "test_scenarios", dataset_root / "test")
    for search_root in search_roots:
        if not search_root.is_dir():
            continue
        for candidate in (search_root, *search_root.rglob("*")):
            if not candidate.is_dir() or candidate == dataset_root / "test":
                continue
            images_dir = candidate / "images"
            labels_dir = candidate / "labels"
            if images_dir.is_dir() and labels_dir.is_dir():
                physical.append({
                    "name": display_scenario(candidate.name),
                    "path": str(candidate.resolve()),
                    "prefix": None,
                })

    if physical:
        unique = {(item["name"].casefold(), item["path"]): item for item in physical}
        return sorted(
            unique.values(),
            key=lambda item: (
                SCENARIO_ORDER.get(item["name"], len(SCENARIO_ORDER)),
                item["name"].casefold(),
            ),
        )

    grouped = {}
    for images_dir in test_directories:
        for image_path, _label_path in _image_label_pairs(images_dir):
            key = _scenario_key_from_file(image_path.name)
            if key:
                grouped.setdefault(key, images_dir)
    return sorted(
        (
            {
                "name": display_scenario(prefix),
                "path": str(images_dir.resolve()),
                "prefix": prefix,
            }
            for prefix, images_dir in grouped.items()
        ),
        key=lambda item: (
            SCENARIO_ORDER.get(item["name"], len(SCENARIO_ORDER)),
            item["name"].casefold(),
        ),
    )


def discover_dataset_scenarios(dataset_path, yaml_path=None):
    """Возвращает редактируемые предложения сценариев без фиксированного набора."""
    dataset_root = normalize_dataset_path(dataset_path)
    yaml_path = Path(yaml_path).expanduser().resolve() if yaml_path else find_dataset_yaml(dataset_root)
    if not yaml_path or not yaml_path.is_file():
        raise ValueError("В выбранном датасете не найден data.yaml.")
    with yaml_path.open("r", encoding="utf-8") as yaml_file:
        data = yaml.safe_load(yaml_file)
    if not isinstance(data, dict):
        raise ValueError("data.yaml должен содержать словарь параметров.")
    test_directories = _resolve_split_directories(dataset_root, yaml_path, data, "test")
    return _auto_scenario_definitions(dataset_root, test_directories)


def _collect_scenarios(dataset_root, test_directories, class_names, definitions=None):
    definitions = (
        _auto_scenario_definitions(dataset_root, test_directories)
        if definitions is None else definitions
    )
    scenarios = []
    used_names = set()
    for index, definition in enumerate(definitions, start=1):
        if isinstance(definition, dict):
            name = str(definition.get("name", "")).strip()
            source_path = definition.get("path", "")
            prefix = definition.get("prefix")
        else:
            name, source_path = definition[:2]
            prefix = definition[2] if len(definition) > 2 else None
            name = str(name).strip()
        if not name:
            raise ValueError(f"Сценарий {index}: не указано название.")
        if name.casefold() in used_names:
            raise ValueError(f"Название сценария «{name}» повторяется.")
        used_names.add(name.casefold())

        images_dir = _as_image_directory(Path(source_path).expanduser().resolve())
        if not images_dir:
            raise ValueError(f"Сценарий «{name}»: не найдена папка изображений: {source_path}")
        pairs = _image_label_pairs(images_dir)
        if prefix:
            prefix = str(prefix).casefold()
            pairs = [
                pair for pair in pairs
                if _scenario_key_from_file(pair[0].name) == prefix
            ]
        stats = _analyze_pairs(pairs, class_names, f"Test / {name}")
        source_text = str(images_dir.parent if images_dir.name.lower() == "images" else images_dir)
        if prefix:
            source_text = f"{images_dir} (префикс: {prefix})"
        stats.update({"name": name, "path": source_text})
        scenarios.append(stats)

    return scenarios


def collect_dataset_statistics(dataset_path, yaml_path=None, scenarios=None):
    dataset_root = normalize_dataset_path(dataset_path)
    yaml_path = Path(yaml_path).expanduser().resolve() if yaml_path else find_dataset_yaml(dataset_root)
    if not yaml_path or not yaml_path.is_file():
        raise ValueError("В выбранном датасете не найден data.yaml.")
    with yaml_path.open("r", encoding="utf-8") as yaml_file:
        data = yaml.safe_load(yaml_file)
    if not isinstance(data, dict):
        raise ValueError("data.yaml должен содержать словарь параметров.")
    class_names = _class_names(data)

    splits = {}
    all_problems = []
    split_directories = {}
    for split in ("train", "val", "test"):
        paths = _resolve_split_directories(dataset_root, yaml_path, data, split)
        split_directories[split] = paths
        parts = [
            _analyze_pairs(_image_label_pairs(path), class_names, SPLIT_LABELS[split])
            for path in paths
        ]
        split_stats = _merge_split_statistics(parts, paths, len(class_names))
        splits[split] = split_stats
        all_problems.extend(split_stats["problems"])

    scenarios = _collect_scenarios(
        dataset_root, split_directories["test"], class_names, scenarios
    )
    return {
        "dataset_root": str(dataset_root),
        "yaml_path": str(yaml_path),
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "class_names": class_names,
        "display_class_names": [display_class(name) for name in class_names],
        "splits": splits,
        "scenarios": scenarios,
        "problems": all_problems,
    }


TITLE_FILL = PatternFill("solid", fgColor="2F80ED")
HEADER_FILL = PatternFill("solid", fgColor="DCEBFA")
SUBTLE_FILL = PatternFill("solid", fgColor="F2F5F9")
WHITE_FONT = Font(color="FFFFFF", bold=True, size=14)
HEADER_FONT = Font(color="243B5A", bold=True)
THIN_BORDER = Border(bottom=Side(style="thin", color="DCE3ED"))


def _style_title(sheet, title, end_column=5):
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    cell = sheet.cell(1, 1, title)
    cell.fill = TITLE_FILL
    cell.font = WHITE_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 28
    sheet.sheet_view.showGridLines = False


def _style_header(row):
    for cell in row:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def _add_table(sheet, ref, name):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    sheet.add_table(table)


def _ordered_class_indices(display_names):
    return sorted(
        range(len(display_names)),
        key=lambda index: (
            CLASS_ORDER.get(display_names[index], len(CLASS_ORDER)), index
        ),
    )


def _configure_reference_chart(
    chart, width=16, height=9.5, y_axis_title=None, show_legend=True
):
    """Оформляет диаграмму как стандартную гистограмму из референса."""
    chart.title = None
    chart.x_axis.title = None
    chart.x_axis.delete = False
    chart.x_axis.tickLblPos = "nextTo"
    chart.x_axis.tickLblSkip = 1
    chart.y_axis.title = y_axis_title
    if chart.y_axis.title is not None:
        for paragraph in chart.y_axis.title.tx.rich.p:
            for run in paragraph.r:
                run.rPr = CharacterProperties(sz=900, b=False)
    chart.y_axis.scaling.min = 0
    chart.y_axis.delete = False
    chart.y_axis.tickLblPos = "nextTo"
    chart.y_axis.majorTickMark = "none"
    chart.y_axis.numFmt = "0"
    chart.y_axis.majorGridlines = ChartLines(
        spPr=GraphicalProperties(
            ln=LineProperties(solidFill="D9D9D9", w=9525)
        )
    )
    if show_legend:
        chart.legend.position = "b"
        chart.legend.overlay = False
    else:
        chart.legend = None
    chart.layout = Layout(
        manualLayout=ManualLayout(
            x=0.07,
            y=0.05,
            w=0.82,
            h=0.72 if show_legend else 0.82,
        )
    )
    chart.width = width
    chart.height = height
    chart.graphicalProperties = GraphicalProperties(
        ln=LineProperties(solidFill="D9D9D9", w=9525)
    )


def _apply_point_colors(series, colors):
    """Задаёт отдельный цвет каждому столбцу одиночного ряда."""
    series.dPt = []
    for index, color in enumerate(colors):
        point = DataPoint(idx=index)
        point.graphicalProperties.solidFill = color
        point.graphicalProperties.line.solidFill = color
        series.dPt.append(point)


def _write_split_sheet(workbook, sheet_name, split_name, stats, class_names, display_names):
    sheet = workbook.create_sheet(sheet_name)
    _style_title(sheet, f"{SPLIT_LABELS[split_name]} выборка", 6)
    metadata = (
        ("Папки изображений", "\n".join(stats["paths"]) or "Не найдены"),
        ("Количество изображений", stats["image_count"]),
        ("Изображений с объектами", stats["annotated_images"]),
        ("Пустых файлов разметки", stats["empty_annotations"]),
        ("Изображений без разметки", stats["missing_labels"]),
        ("Некорректных строк", stats["invalid_lines"]),
    )
    for row_index, (label, value) in enumerate(metadata, start=3):
        sheet.cell(row_index, 1, label).font = HEADER_FONT
        sheet.cell(row_index, 2, value)
        sheet.cell(row_index, 1).fill = SUBTLE_FILL
    sheet.merge_cells("B3:F3")
    sheet["B3"].alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[3].height = 32
    header_row = 11
    headers = ("ID", "Класс", "Имя в data.yaml", "Количество объектов", "Изображений с классом", "Доля объектов")
    for column, value in enumerate(headers, start=1):
        sheet.cell(header_row, column, value)
    _style_header(sheet[header_row])
    total_objects = sum(stats["instances"])
    class_order = _ordered_class_indices(display_names)
    for offset, class_id in enumerate(class_order, start=1):
        raw_name = class_names[class_id]
        shown_name = display_names[class_id]
        row = header_row + offset
        sheet.cell(row, 1, class_id)
        sheet.cell(row, 2, shown_name)
        sheet.cell(row, 3, raw_name)
        sheet.cell(row, 4, stats["instances"][class_id])
        sheet.cell(row, 5, stats["images_with_class"][class_id])
        sheet.cell(row, 6, f"=IF(SUM($D${header_row + 1}:$D${header_row + len(class_names)})=0,0,D{row}/SUM($D${header_row + 1}:$D${header_row + len(class_names)}))")
        sheet.cell(row, 6).number_format = "0.0%"
    end_row = header_row + len(class_names)
    _add_table(sheet, f"A{header_row}:F{end_row}", f"{sheet_name}ClassDistribution")

    chart = BarChart()
    chart.type = "col"
    chart.add_data(Reference(sheet, min_col=4, min_row=header_row, max_row=end_row), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=2, min_row=header_row + 1, max_row=end_row))
    chart.series[0].tx = SeriesLabel(v="Количество объектов")
    chart.varyColors = True
    _apply_point_colors(
        chart.series[0],
        [EXCEL_SERIES_COLORS.get(display_names[class_id], "4F81BD") for class_id in class_order],
    )
    _configure_reference_chart(
        chart,
        y_axis_title="Количество экземпляров",
        show_legend=False,
    )
    sheet.add_chart(chart, "H2")

    sheet.freeze_panes = f"A{header_row + 1}"
    widths = {"A": 27, "B": 28, "C": 22, "D": 22, "E": 23, "F": 17, "G": 3}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    return total_objects


def write_dataset_statistics(stats, output_path):
    output_path = Path(output_path).expanduser().resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Сводка"
    _style_title(summary, "Статистика датасета", 5)
    summary["A3"] = "Корень датасета"
    summary["B3"] = stats["dataset_root"]
    summary["A4"] = "data.yaml"
    summary["B4"] = stats["yaml_path"]
    summary["A5"] = "Сформировано"
    summary["B5"] = stats["generated_at"]
    for cell in (summary["A3"], summary["A4"], summary["A5"]):
        cell.font = HEADER_FONT
        cell.fill = SUBTLE_FILL

    split_sheet_names = {"train": "Train", "val": "Val", "test": "Test"}
    for split, sheet_name in split_sheet_names.items():
        _write_split_sheet(
            workbook, sheet_name, split, stats["splits"][split],
            stats["class_names"], stats["display_class_names"],
        )

    summary_headers = ("Выборка", "Количество изображений", "Изображений с объектами", "Без разметки", "Пустая разметка")
    for column, value in enumerate(summary_headers, start=1):
        summary.cell(8, column, value)
    _style_header(summary[8])
    for row, (split, sheet_name) in enumerate(split_sheet_names.items(), start=9):
        summary.cell(row, 1, SPLIT_LABELS[split])
        summary.cell(row, 2, f"='{sheet_name}'!B4")
        summary.cell(row, 3, f"='{sheet_name}'!B5")
        summary.cell(row, 4, f"='{sheet_name}'!B7")
        summary.cell(row, 5, f"='{sheet_name}'!B6")
    _add_table(summary, "A8:E11", "DatasetSplitSummary")

    split_chart = BarChart()
    split_chart.type = "col"
    split_chart.add_data(Reference(summary, min_col=2, min_row=8, max_row=11), titles_from_data=True)
    split_chart.set_categories(Reference(summary, min_col=1, min_row=9, max_row=11))
    split_chart.series[0].tx = SeriesLabel(v="Количество изображений")
    split_chart.varyColors = True
    _apply_point_colors(split_chart.series[0], SUMMARY_POINT_COLORS)
    _configure_reference_chart(
        split_chart,
        width=14,
        height=8.5,
        y_axis_title="Количество изображений",
        show_legend=False,
    )
    summary.add_chart(split_chart, "G2")
    summary.column_dimensions["A"].width = 27
    summary.column_dimensions["B"].width = 66
    for column in ("C", "D", "E"):
        summary.column_dimensions[column].width = 23
    summary["B3"].alignment = Alignment(wrap_text=True)
    summary["B4"].alignment = Alignment(wrap_text=True)
    summary.freeze_panes = "A8"

    scenarios_sheet = workbook.create_sheet("Сценарии test")
    scenario_end_column = max(3, len(stats["class_names"]) + 2)
    _style_title(scenarios_sheet, "Распределение тестовой выборки по сценариям", scenario_end_column)
    scenarios_sheet["A3"] = "Источник сценариев"
    scenarios_sheet["B3"] = (
        "Папки сценариев или префиксы имен файлов"
        if stats["scenarios"] else "Сценарии не обнаружены"
    )
    scenarios_sheet["A3"].font = HEADER_FONT
    scenarios_sheet["A3"].fill = SUBTLE_FILL
    scenario_class_order = _ordered_class_indices(stats["display_class_names"])
    ordered_scenario_classes = [
        stats["display_class_names"][index] for index in scenario_class_order
    ]
    headers = ["Сценарий", "Изображений", *ordered_scenario_classes]
    for column, value in enumerate(headers, start=1):
        scenarios_sheet.cell(6, column, value)
    _style_header(scenarios_sheet[6])
    for row, scenario in enumerate(stats["scenarios"], start=7):
        scenarios_sheet.cell(row, 1, scenario["name"])
        scenarios_sheet.cell(row, 2, scenario["image_count"])
        for column, class_index in enumerate(scenario_class_order, start=3):
            scenarios_sheet.cell(row, column, scenario["instances"][class_index])
    if stats["scenarios"]:
        end_row = 6 + len(stats["scenarios"])
        _add_table(scenarios_sheet, f"A6:{scenarios_sheet.cell(6, scenario_end_column).column_letter}{end_row}", "TestScenarioDistribution")
        scenario_chart = BarChart()
        scenario_chart.type = "col"
        scenario_chart.grouping = "clustered"
        scenario_chart.add_data(
            Reference(scenarios_sheet, min_col=3, max_col=scenario_end_column, min_row=6, max_row=end_row),
            titles_from_data=True,
        )
        scenario_chart.set_categories(Reference(scenarios_sheet, min_col=1, min_row=7, max_row=end_row))
        for series, class_name in zip(scenario_chart.series, ordered_scenario_classes):
            series.tx = SeriesLabel(v=class_name)
            color = EXCEL_SERIES_COLORS.get(class_name, "4F81BD")
            series.graphicalProperties.solidFill = color
            series.graphicalProperties.line.solidFill = color
        _configure_reference_chart(
            scenario_chart,
            width=16,
            height=9.5,
            y_axis_title="Количество экземпляров",
        )
        scenarios_sheet.add_chart(scenario_chart, "H2")
    scenarios_sheet.column_dimensions["A"].width = 22
    scenarios_sheet.column_dimensions["B"].width = 17
    for column in range(3, scenario_end_column + 1):
        scenarios_sheet.column_dimensions[get_column_letter(column)].width = 18
    scenarios_sheet.freeze_panes = "A7"

    problems_sheet = workbook.create_sheet("Проблемы")
    _style_title(problems_sheet, "Проверка целостности разметки", 5)
    problem_headers = ("Раздел", "Изображение", "Файл разметки", "Проблема", "Подробности")
    for column, value in enumerate(problem_headers, start=1):
        problems_sheet.cell(3, column, value)
    _style_header(problems_sheet[3])
    if stats["problems"]:
        for row, problem in enumerate(stats["problems"], start=4):
            problems_sheet.cell(row, 1, problem["section"])
            problems_sheet.cell(row, 2, problem["image"])
            problems_sheet.cell(row, 3, problem["label"])
            problems_sheet.cell(row, 4, problem["problem"])
            problems_sheet.cell(row, 5, problem["details"])
        _add_table(problems_sheet, f"A3:E{3 + len(stats['problems'])}", "DatasetProblems")
    else:
        problems_sheet["A4"] = "Проблем не обнаружено"
    for column, width in {"A": 20, "B": 58, "C": 58, "D": 30, "E": 46}.items():
        problems_sheet.column_dimensions[column].width = width
    problems_sheet.freeze_panes = "A4"

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(output_path)
    return output_path


def generate_dataset_statistics_report(
    dataset_path, output_path=None, yaml_path=None, scenarios=None
):
    stats = collect_dataset_statistics(
        dataset_path, yaml_path=yaml_path, scenarios=scenarios
    )
    output_path = output_path or Path(stats["dataset_root"]) / "Статистика_датасета.xlsx"
    return stats, write_dataset_statistics(stats, output_path)
